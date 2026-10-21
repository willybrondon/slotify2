# -*- coding: utf-8 -*-
"""Generate Skedisy Facebook group posts spreadsheet (Client(e) + Pro)."""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path

out = Path(__file__).with_name("SKEDISY_POSTS_FACEBOOK_GROUPES.xlsx")

groups = [
    ("Black Geek / Recherche Coiffeuse Afro Pro IDF", "https://www.facebook.com/groups/134449370024416"),
    ("Coiffeuses Afro professionnelle en île de France", "https://www.facebook.com/groups/284571943882019"),
    ("Coiffeuses Afro Certifiées & Professionnelles IDF", "https://www.facebook.com/groups/1142079080008153"),
    ("JE CONSOMME AFRICAIN", "https://www.facebook.com/groups/2213701955383712/"),
    ("La Diaspora Africaine de France", "https://www.facebook.com/groups/1539354113070702"),
    ("Coiffeuse à domicile 93", "https://www.facebook.com/groups/4769417296405741"),
    ("Soins, Astuces et Conseils pour Cheveux Naturels", "https://www.facebook.com/groups/5749916125108167"),
    ("La femme belle en coiffure", "https://www.facebook.com/groups/667615722313297"),
    ("Coiffure, top modèle 2026", "https://www.facebook.com/groups/1584477032567635"),
    ("Coiffure afro à domicile France", "https://www.facebook.com/groups/1240347616412450"),
    ("Coiffeuses afro en île de France et ailleurs", "https://www.facebook.com/groups/1948257189237541/"),
    ("Glamour et gloire coiffure afro île de France", "https://www.facebook.com/groups/970050834928729/"),
    ("coiffeuse Afro pro 72", "https://www.facebook.com/groups/994867493011335/"),
    ("coiffeuse afros a domicile", "https://www.facebook.com/groups/767740608226140/"),
    ("Les locksés d’île de France", "https://www.facebook.com/groups/2692720617599643/"),
    ("Votre coiffeuse Afro en Île-de-France", "https://www.facebook.com/groups/1331027702017734/"),
    ("Tresses afro en Île-de-France", "https://www.facebook.com/groups/241358959612399/"),
    ("Coiffure afro africaine île de France", "https://www.facebook.com/groups/259213433910722/"),
    ("Coiffure à domicile à Paris et île de france", "https://www.facebook.com/groups/1347267503122986/"),
    ("coiffures Afro au naturel", "https://www.facebook.com/groups/1324886558019993/"),
    ("coiffures Afro au naturel (2)", "https://www.facebook.com/groups/900225391929163/"),
    ("Coiffeuses Afro en Île de France", "https://www.facebook.com/groups/692341087084067/"),
    ("SOS BEAUTÉS AFRO", "https://www.facebook.com/groups/105919819792897/"),
    ("Recherche Coiffeuse Afro Pro Ile-de-France (2)", "https://www.facebook.com/groups/333981471660575/"),
    ("Coiffeuse afro professionnelle", "https://www.facebook.com/groups/1236182034790135/"),
    ("Astuces beauté, conseils soin de la peau et cheveux", "https://www.facebook.com/groups/1249795580163427/"),
    ("Recettes beautés pour tout type de peaux", "https://www.facebook.com/groups/3016520931749842/"),
    ("Comment s'habiller dans un rendez-vous", "https://www.facebook.com/groups/462382094844180/"),
    ("Coiffure à domicile 95,94,93,92,91,78,77 et 27", "https://www.facebook.com/groups/302250931646026/"),
]

client_posts = [
    {
        "groupe": groups[0][0],
        "lien": groups[0][1],
        "angle": "Chaos WhatsApp / réserver simplement",
        "post": (
            "On a arrêté de chercher une coiffeuse afro en Île-de-France dans 15 conversations WhatsApp.\n\n"
            "Avec Skedisy, tu vois les salons, les prestations (tresses, knotless, locks…) et tu réserves un vrai créneau — sans négocier toute la soirée.\n\n"
            "Moins de « tu es dispo samedi ? », plus de rendez-vous confirmé. On a trouvé le moyen de transformer le chaos des messages en réservation claire.\n\n"
            "Si toi aussi tu en as marre de courir après une place : découvre Skedisy et réserve sur skedisy.com (app gratuite)."
        ),
    },
    {
        "groupe": groups[3][0],
        "lien": groups[3][1],
        "angle": "Consommer africain / local IDF",
        "post": (
            "Consommer africain, ce n’est pas seulement acheter un produit — c’est aussi soutenir les salons afro près de chez toi.\n\n"
            "Skedisy met en avant les salons de beauté afro en Île-de-France : tu trouves, tu compares, tu réserves en quelques clics.\n\n"
            "Derrière chaque coiffure, il y a une experte, un savoir-faire, une communauté. On a construit la plateforme pour que cette valeur soit enfin visible et bookable.\n\n"
            "Partage Skedisy autour de toi et réserve ton prochain look sur skedisy.com."
        ),
    },
    {
        "groupe": groups[4][0],
        "lien": groups[4][1],
        "angle": "Diaspora / appartenance",
        "post": (
            "La diaspora connaît déjà le vrai problème : trouver LA bonne coiffeuse afro, sans mauvaises surprises.\n\n"
            "Skedisy te donne une vitrine claire des salons en IDF — prestations, avis, créneaux — pour réserver sereinement.\n\n"
            "Fini le hasard du bouche-à-oreille seul : on a relié la demande de la communauté à des professionnelles prêtes à t’accueillir.\n\n"
            "Rejoins le mouvement : teste Skedisy et envoie le lien à une sœur qui cherche encore."
        ),
    },
    {
        "groupe": groups[6][0],
        "lien": groups[6][1],
        "angle": "Cheveux naturels / conseils",
        "post": (
            "On parle soins et cheveux naturels tous les jours… mais le RDV, lui, se joue encore en DM.\n\n"
            "Skedisy t’aide à trouver un salon adapté à tes cheveux texturés en Île-de-France et à bloquer un créneau sans friction.\n\n"
            "Astuces + bon salon + date claire = moins de stress. On a relié le conseil à l’action : la réservation.\n\n"
            "Garde ce post, ouvre skedisy.com, et book ton prochain soin / coiffure."
        ),
    },
    {
        "groupe": groups[7][0],
        "lien": groups[7][1],
        "angle": "Beauté / confiance",
        "post": (
            "Se sentir belle commence souvent… avant d’arriver au salon.\n\n"
            "Avec Skedisy, tu choisis ta presta afro (tresses, tissage, locks…) et tu arrives avec un rendez-vous déjà posé.\n\n"
            "Quand le créneau est clair, la tête est légère. On a conçu ça pour que la beauté ne passe plus par l’angoisse du « est-ce qu’elles vont me répondre ? ».\n\n"
            "Fais-toi ce cadeau : réserve sur Skedisy et partage à une amie."
        ),
    },
    {
        "groupe": groups[8][0],
        "lien": groups[8][1],
        "angle": "Tendance 2026",
        "post": (
            "Les looks 2026 arrivent. La vraie question : qui peut te les faire… samedi ?\n\n"
            "Skedisy te montre les salons afro en IDF et les créneaux dispo pour booker le style que tu veux.\n\n"
            "Tendance sans place = frustration. Tendance + réservation = résultat. On a trouvé le pont entre l’inspiration et le fauteuil.\n\n"
            "Enregistre le look, ouvre Skedisy, et réserve avant que le créneau parte."
        ),
    },
    {
        "groupe": groups[14][0],
        "lien": groups[14][1],
        "angle": "Locks IDF",
        "post": (
            "Les locks, ça ne se confie pas au premier salon venu.\n\n"
            "Sur Skedisy, tu cherches des professionnelles qui comprennent vraiment les locks en Île-de-France — et tu réserves un créneau adapté.\n\n"
            "Entretien, reprise, installation : le détail compte. On a bâti une plateforme où cette expertise n’est plus noyée dans « coiffure générale ».\n\n"
            "Si tu es locksé·e (ou tu veux le devenir) : trouve ton salon sur skedisy.com et book."
        ),
    },
    {
        "groupe": groups[16][0],
        "lien": groups[16][1],
        "angle": "Tresses IDF",
        "post": (
            "« Qui fait des tresses propres près de chez moi ? » — le message qu’on voit trop souvent.\n\n"
            "Skedisy répond sans le groupe : salons afro IDF, prestations visibles, réservation en ligne.\n\n"
            "Tu ne quémandes plus une place : tu choisis. On a remplacé la quête dans les commentaires par un agenda réel.\n\n"
            "Réserve directement sur Skedisy et tague une amie qui cherche encore."
        ),
    },
    {
        "groupe": groups[5][0],
        "lien": groups[5][1],
        "angle": "Domicile 93",
        "post": (
            "Le 93 bouge. Les coiffeuses aussi — domicile ou salon, le besoin est le même : un RDV clair.\n\n"
            "Skedisy centralise des professionnelles beauté afro en Île-de-France pour que tu réserves sans enchaîner les appels.\n\n"
            "Quand le créneau est posé, tout le monde gagne du temps. On a simplifié ce que WhatsApp compliquait.\n\n"
            "Habitante du 93 (ou d’à côté) : jette un œil sur skedisy.com et book ton prochain rendez-vous."
        ),
    },
    {
        "groupe": groups[9][0],
        "lien": groups[9][1],
        "angle": "Domicile France",
        "post": (
            "Coiffure afro à domicile, c’est pratique — jusqu’au moment où personne ne confirme.\n\n"
            "Skedisy t’aide à trouver et réserver auprès de professionnelles en Île-de-France, avec un parcours clair.\n\n"
            "Moins de « je te dirai demain », plus de date dans l’agenda. On a mis de l’ordre dans la prise de RDV.\n\n"
            "Envoie ce post à celle qui cherche encore, et réserve sur Skedisy."
        ),
    },
    {
        "groupe": groups[18][0],
        "lien": groups[18][1],
        "angle": "Domicile Paris IDF",
        "post": (
            "Paris et IDF : beaucoup d’options, trop peu de clarté.\n\n"
            "Skedisy filtre pour la beauté afro : tu vois, tu choisis, tu réserves — salon ou parcours pro local.\n\n"
            "La ville va vite ; ton RDV ne doit pas traîner 10 jours en messages. On a accéléré le passage à l’action.\n\n"
            "Prochain week-end libre ? Bloque-le sur skedisy.com maintenant."
        ),
    },
    {
        "groupe": groups[19][0],
        "lien": groups[19][1],
        "angle": "Afro au naturel",
        "post": (
            "Cheveux afro au naturel, on veut du respect — pas une coupe « au feeling ».\n\n"
            "Skedisy te met face à des salons qui comprennent texture, volume, et vrai projet coiffure en IDF.\n\n"
            "Tu arrives préparée, avec une presta claire. On a rendu visible ce que les agendas généralistes oublient.\n\n"
            "Pour ton prochain naturel / twist-out / soin : passe par Skedisy et partage le lien."
        ),
    },
    {
        "groupe": groups[22][0],
        "lien": groups[22][1],
        "angle": "SOS beautés afro",
        "post": (
            "SOS beauté afro : urgence de RDV, pas d’urgence de stress.\n\n"
            "Skedisy te montre qui peut te prendre en Île-de-France et te laisse réserver sans négocier 40 messages.\n\n"
            "Quand tu as besoin d’être sauvée pour un événement, chaque heure compte. On a construit le raccourci.\n\n"
            "Sauvegarde Skedisy, réserve, et envoie le secours à une amie en panique."
        ),
    },
    {
        "groupe": groups[25][0],
        "lien": groups[25][1],
        "angle": "Glow / peaux & cheveux",
        "post": (
            "Le glow, c’est la peau… et aussi une coiffure qui tient le coup.\n\n"
            "Avec Skedisy, tu bookes un salon afro en IDF comme tu bookerais un vrai plan beauté : date, presta, confirmation.\n\n"
            "Routine + rendez-vous = résultat. On a relié le conseil « glow » à un créneau concret.\n\n"
            "Ajoute Skedisy à ta routine : skedisy.com — et tague ta glow partner."
        ),
    },
    {
        "groupe": groups[26][0],
        "lien": groups[26][1],
        "angle": "Recettes beauté / peaux",
        "post": (
            "On partage des recettes beauté… et on oublie parfois de booker le pro qui termine le travail.\n\n"
            "Skedisy complète ta routine DIY : trouve un salon afro en Île-de-France et réserve en ligne.\n\n"
            "Maison + salon = combo gagnant. On a rendu le second aussi simple que le premier.\n\n"
            "Garde ta recette, et réserve le RDV sur Skedisy."
        ),
    },
    {
        "groupe": groups[27][0],
        "lien": groups[27][1],
        "angle": "Look pour un rendez-vous",
        "post": (
            "Tu sais quoi porter pour le rendez-vous. Et la coiffure ?\n\n"
            "Skedisy t’aide à bloquer un créneau chez une pro afro en IDF avant le jour J — sans improvisation.\n\n"
            "Une tenue forte + des cheveux prêts = confiance. On a enlevé le stress de la « dernière minute ».\n\n"
            "Book ton look capillaire sur skedisy.com et envoie ce post à celle qui a un date bientôt."
        ),
    },
    {
        "groupe": groups[11][0],
        "lien": groups[11][1],
        "angle": "Glamour afro IDF",
        "post": (
            "Le glamour afro en Île-de-France mérite mieux que le hasard.\n\n"
            "Skedisy te donne accès à des salons spécialisés : tu choisis la presta, tu réserves, tu arrives sereine.\n\n"
            "Quand la plateforme parle ton langage (tresses, lace, locks…), tout change. On a créé cet espace pour ça.\n\n"
            "Passe en mode glamour : réserve sur Skedisy et partage."
        ),
    },
    {
        "groupe": groups[17][0],
        "lien": groups[17][1],
        "angle": "Coiffure afro africaine IDF",
        "post": (
            "La coiffure afro africaine, c’est un héritage — pas une case « autres » dans un agenda générique.\n\n"
            "Skedisy est pensé pour la beauté afro en Île-de-France : recherche, fiche salon, réservation.\n\n"
            "Tu vois enfin des professionnelles alignées avec ce que tu cherches vraiment. On a refusé le flou.\n\n"
            "Honore ton style : book sur skedisy.com et fais tourner dans le groupe."
        ),
    },
    {
        "groupe": groups[23][0],
        "lien": groups[23][1],
        "angle": "Recherche coiffeuse",
        "post": (
            "On a tous déjà posté « je cherche une coiffeuse afro sérieuse en IDF ».\n\n"
            "Skedisy transforme cette recherche en résultats + créneaux réservables.\n\n"
            "Le post dans le groupe lance la conversation ; la plateforme termine le job. On a trouvé la suite logique.\n\n"
            "Arrête de republier la même demande : ouvre Skedisy et réserve."
        ),
    },
    {
        "groupe": groups[28][0],
        "lien": groups[28][1],
        "angle": "Domicile multi-départements",
        "post": (
            "95, 94, 93, 92, 91, 78, 77… le besoin est le même partout : une coiffeuse afro fiable, un créneau clair.\n\n"
            "Skedisy couvre l’Île-de-France pour que tu réserves sans cartographier tous les groupes Facebook.\n\n"
            "Un seul réflexe pour plusieurs départements. On a centralisé ce que tu cherchais en silos.\n\n"
            "Quel que soit ton code postal : tente skedisy.com et partage à ton cercle."
        ),
    },
]

pro_posts = [
    {
        "groupe": groups[1][0],
        "lien": groups[1][1],
        "angle": "Agenda / moins de DM",
        "post": (
            "On a arrêté de gérer tout le salon dans Instagram et WhatsApp.\n\n"
            "Avec Skedisy, tes clientes en Île-de-France voient tes prestas afro et réservent — tu reçois des RDV préparés, pas 40 messages à minuit.\n\n"
            "Moins de chaos, plus de fauteuils remplis. On a construit l’outil pour les salons qui font vraiment du knotless, des tresses, des locks.\n\n"
            "Rejoins gratuitement sur skedisy.com/professionnel et envoie ce post à une collègue."
        ),
    },
    {
        "groupe": groups[2][0],
        "lien": groups[2][1],
        "angle": "Certifiées / image pro",
        "post": (
            "Être certifiée et pro, ça se voit — sauf quand ta seule vitrine, c’est un fil Instagram saturé.\n\n"
            "Skedisy te donne une fiche salon claire : prestations, créneaux, réservations — pensée beauté afro IDF.\n\n"
            "Tu ne te bats plus seulement contre l’algo : tu es trouvable quand une cliente cherche. On a ouvert ce canal.\n\n"
            "Crée ton espace (gratuit pour commencer) sur Skedisy et partage dans le groupe."
        ),
    },
    {
        "groupe": groups[10][0],
        "lien": groups[10][1],
        "angle": "Visibilité IDF et ailleurs",
        "post": (
            "Tu travailles en IDF (et parfois ailleurs) — mais les clientes te trouvent encore au hasard.\n\n"
            "Skedisy te met devant celles qui cherchent déjà une coiffeuse afro près de chez elles.\n\n"
            "Visibilité ciblée > posts perdus dans le feed. On a relié la demande réelle à ton agenda.\n\n"
            "Inscris ton salon sur skedisy.com/professionnel et tague une pro du groupe."
        ),
    },
    {
        "groupe": groups[12][0],
        "lien": groups[12][1],
        "angle": "Pro hors Paris / digital",
        "post": (
            "Pas besoin d’être à Paris pour digitaliser ton activité afro.\n\n"
            "Skedisy aide les professionnelles à structurer la réservation et à être trouvées par des clientes qui cherchent du sérieux.\n\n"
            "Le pro qui gagne, c’est celui dont le créneau est clair. On a rendu ça simple.\n\n"
            "Découvre Skedisy Pro et partage la solution avec le groupe."
        ),
    },
    {
        "groupe": groups[13][0],
        "lien": groups[13][1],
        "angle": "Domicile — pro",
        "post": (
            "À domicile, le téléphone sonne… et le planning se casse.\n\n"
            "Skedisy t’aide à poser des rendez-vous clairs pour tes clientes afro, sans vivre dans les messages.\n\n"
            "Moins de no-shows flous, plus de journées organisées. On a adapté la réservation à ton métier.\n\n"
            "Passe pro sur Skedisy (gratuit pour commencer) et envoie le lien à une collègue à domicile."
        ),
    },
    {
        "groupe": groups[15][0],
        "lien": groups[15][1],
        "angle": "Votre coiffeuse — acquisition",
        "post": (
            "« Votre coiffeuse afro en IDF » — le titre du groupe. Et si les clientes te trouvaient vraiment ?\n\n"
            "Skedisy apporte des demandes structurées : presta, date, réservation — pas seulement des likes.\n\n"
            "Acquisition + agenda = croissance. On a relié les deux pour les salons afro.\n\n"
            "Rejoins Skedisy et partage ce post aux pros du groupe."
        ),
    },
    {
        "groupe": groups[21][0],
        "lien": groups[21][1],
        "angle": "Coiffeuses Afro IDF",
        "post": (
            "Les coiffeuses afro d’Île-de-France méritent un outil à leur image — pas un agenda « coupe femme » générique.\n\n"
            "Skedisy est spécialisé beauté afro : tes techniques existent, tes clientes te trouvent, elles réservent.\n\n"
            "Spécialisation = moins de concurrence inutile, plus de leads pertinents. On a choisi ce positionnement pour toi.\n\n"
            "Ouvre ton compte salon sur skedisy.com/professionnel — c’est gratuit pour commencer."
        ),
    },
    {
        "groupe": groups[24][0],
        "lien": groups[24][1],
        "angle": "Professionnelle — filling slots",
        "post": (
            "Être pro, c’est aussi remplir les trous du planning sans brader ton prix dans la rue.\n\n"
            "Sur Skedisy, tu ouvres des créneaux et tu reçois des clientes déjà décidées en IDF.\n\n"
            "Tu restes maîtresse de ton offre ; la plateforme amène la demande. On a trouvé cet équilibre.\n\n"
            "Active Skedisy et partage la méthode dans le groupe."
        ),
    },
    {
        "groupe": groups[0][0],
        "lien": groups[0][1],
        "angle": "Recherche coiffeuse = opportunité pro",
        "post": (
            "Chaque jour, des clientes postent « je cherche une coiffeuse afro pro en IDF ».\n\n"
            "Skedisy transforme ces recherches en réservations sur ta fiche — tu n’as plus à répondre à tout le monde en DM.\n\n"
            "Le groupe révèle le besoin ; la plateforme le convertit. On a construit le pont.\n\n"
            "Si tu es la solution : inscris-toi sur Skedisy et envoie ce post."
        ),
    },
    {
        "groupe": groups[11][0],
        "lien": groups[11][1],
        "angle": "Glamour — image salon",
        "post": (
            "Le glamour se voit sur les photos — le business se voit sur l’agenda rempli.\n\n"
            "Skedisy te donne les deux : vitrine + réservations pour salons afro en Île-de-France.\n\n"
            "Beauté sans booking, c’est du contenu. Beauté + booking, c’est du CA. On a relié les deux.\n\n"
            "Passe sur Skedisy Pro et partage avec les salons du groupe."
        ),
    },
    {
        "groupe": groups[14][0],
        "lien": groups[14][1],
        "angle": "Locks — expertise pro",
        "post": (
            "Spécialiste locks ? Les clientes te cherchent — encore faut-il qu’elles te bookent.\n\n"
            "Skedisy met en avant les expertises afro (locks, tresses, etc.) pour l’Île-de-France.\n\n"
            "Ton savoir-faire mérite mieux qu’un message « dispo ? ». On a rendu l’expertise réservable.\n\n"
            "Crée ta fiche Skedisy et invite les lockseuses du groupe."
        ),
    },
    {
        "groupe": groups[16][0],
        "lien": groups[16][1],
        "angle": "Tresses — volume RDV",
        "post": (
            "Les tresses, c’est du temps. Chaque no-show coûte cher.\n\n"
            "Avec Skedisy, la cliente réserve un vrai créneau : tu organises mieux tes journées de tressage en IDF.\n\n"
            "Moins de flou, plus de prévisibilité. On a conçu ça pour les prestas longues et complexes.\n\n"
            "Rejoins Skedisy gratuitement et partage aux tresseuses du groupe."
        ),
    },
    {
        "groupe": groups[5][0],
        "lien": groups[5][1],
        "angle": "Domicile 93 — pro",
        "post": (
            "Le 93 est concurrentiel. Ce qui change la donne : être bookable en 2 minutes.\n\n"
            "Skedisy te rend visible auprès des clientes afro d’Île-de-France qui veulent réserver sans négocier.\n\n"
            "Le premier qui structure son agenda gagne. On t’apporte l’outil.\n\n"
            "Inscris-toi sur skedisy.com/professionnel et fais tourner."
        ),
    },
    {
        "groupe": groups[9][0],
        "lien": groups[9][1],
        "angle": "Domicile France — pro",
        "post": (
            "À domicile en France, la logistique mange ton énergie.\n\n"
            "Skedisy clarifie la prise de RDV pour que tu gardes du temps pour coiffer — pas pour chasser les confirmations.\n\n"
            "Automatiser la réservation, c’est du temps rendu. On a trouvé ce levier pour les pros afro.\n\n"
            "Passe sur Skedisy et partage la solution aux coiffeuses à domicile."
        ),
    },
    {
        "groupe": groups[18][0],
        "lien": groups[18][1],
        "angle": "Paris IDF — pro",
        "post": (
            "Paris et IDF : trop de bruits, trop peu de leads qualifiés.\n\n"
            "Skedisy filtre : clientes qui cherchent la beauté afro, pas « n’importe quel salon ».\n\n"
            "Qualité de demande > quantité de DM. On a bâti la plateforme autour de ça.\n\n"
            "Rejoins Skedisy Pro et envoie ce post aux salons du groupe."
        ),
    },
    {
        "groupe": groups[28][0],
        "lien": groups[28][1],
        "angle": "Multi-départements — pro",
        "post": (
            "Tu couvres plusieurs départements ? Tes clientes, elles, cherchent « près de moi ».\n\n"
            "Skedisy te place sur la carte beauté afro d’Île-de-France avec réservation en ligne.\n\n"
            "Un profil clair bat dix annonces éparpillées. On a centralisé ton acquisition.\n\n"
            "Active Skedisy et partage aux pros des 95/94/93/92…"
        ),
    },
    {
        "groupe": groups[22][0],
        "lien": groups[22][1],
        "angle": "SOS — répondre à l’urgence",
        "post": (
            "Les SOS beauté afro = des clientes prêtes à booker… si tu es joignable.\n\n"
            "Skedisy te rend bookable 24h/24 : elles réservent même pendant que tu coiffes.\n\n"
            "L’urgence devient un créneau, pas un stress. On a transformé le chaos en pipeline.\n\n"
            "Inscris ton salon sur Skedisy et sauve aussi ton planning."
        ),
    },
    {
        "groupe": groups[3][0],
        "lien": groups[3][1],
        "angle": "Consommer africain — B2B",
        "post": (
            "« Je consomme africain » côté clientes. Côté salons : sois trouvable.\n\n"
            "Skedisy est le canal digital pour les professionnelles afro en IDF qui veulent des réservations, pas seulement de la visibilité vanity.\n\n"
            "Communauté + outil = business. On a relié les deux.\n\n"
            "Rejoins Skedisy et partage dans les groupes qui comptent."
        ),
    },
    {
        "groupe": groups[4][0],
        "lien": groups[4][1],
        "angle": "Diaspora — pro",
        "post": (
            "La diaspora cherche des mains expertes. Encore faut-il un chemin jusqu’à ton fauteuil.\n\n"
            "Skedisy crée ce chemin : découverte + réservation pour salons afro en Île-de-France.\n\n"
            "Tu gardes ton identité ; tu gagnes en structure. On a refusé le modèle généraliste.\n\n"
            "Ouvre ton compte Skedisy Pro et invite une collègue de la diaspora."
        ),
    },
    {
        "groupe": groups[8][0],
        "lien": groups[8][1],
        "angle": "Tendance 2026 — pro",
        "post": (
            "Les tendances 2026 font rêver. Les agendas vides, non.\n\n"
            "Avec Skedisy, tu transformes l’inspiration Instagram en créneaux réservés sur ta fiche salon afro.\n\n"
            "Le contenu attire ; la réservation encaisse. On a fermé la boucle.\n\n"
            "Passe pro sur Skedisy (gratuit pour commencer) et partage ce levier."
        ),
    },
]

header_fill = PatternFill("solid", fgColor="C45C26")
header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
wrap = Alignment(wrap_text=True, vertical="top")
thin = Border(
    left=Side(style="thin", color="DDDDDD"),
    right=Side(style="thin", color="DDDDDD"),
    top=Side(style="thin", color="DDDDDD"),
    bottom=Side(style="thin", color="DDDDDD"),
)


def write_sheet(ws, title, rows):
    ws.title = title
    headers = [
        "N°",
        "Groupe Facebook",
        "Lien du groupe",
        "Angle / thématique",
        "Post (FR) — prêt à publier",
        "Audience",
        "Langue",
        "Plateforme",
    ]
    ws.append(headers)
    for col, _h in enumerate(headers, 1):
        cell = ws.cell(1, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(rows, 1):
        ws.append(
            [
                i,
                r["groupe"],
                r["lien"],
                r["angle"],
                r["post"],
                title,
                "FR",
                "Facebook Group",
            ]
        )
    widths = [6, 42, 48, 28, 80, 14, 10, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 28
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=8):
        for cell in row:
            cell.alignment = wrap
            cell.border = thin
        ws.row_dimensions[row[0].row].height = 120
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{ws.max_row}"


wb = Workbook()
write_sheet(wb.active, "Client(e)", client_posts)
write_sheet(wb.create_sheet(), "Pro", pro_posts)

ws3 = wb.create_sheet("Guide_utilisation")
ws3.append(["Comment utiliser ce fichier"])
ws3.append([])
ws3.append(["Structure de chaque post (4 temps, sans libellés dans le texte publié)"])
ws3.append(["1. Hook viral"])
ws3.append(["2. Proposition de valeur Skedisy"])
ws3.append(["3. Intrigue / amplification"])
ws3.append(["4. Appel à l’action + partage"])
ws3.append([])
ws3.append(["Onglet Client(e) : 20 posts FR clientes, calés sur des groupes Facebook"])
ws3.append(["Onglet Pro : 20 posts FR coiffeuses / salons, calés sur des groupes Facebook"])
ws3.append([])
ws3.append(
    [
        "Pour scaler (ex. 30 groupes × 20 posts) : dupliquer une ligne, changer nom/lien du groupe, adapter 1–2 mots du hook à la thématique."
    ]
)
ws3.append(
    [
        "Skedisy : beauté afro Île-de-France — réservation clientes gratuite — salon gratuit pour commencer — skedisy.com"
    ]
)
ws3.append([])
ws3.append(
    [
        "Importer dans Google Sheets : Fichier > Importer > Téléverser ce .xlsx > Créer de nouvelles feuilles / Ajouter aux feuilles existantes"
    ]
)
ws3.append(
    [
        "Lien modèle : https://docs.google.com/spreadsheets/d/1MzNMdF2jnIMCugWvYpbcVOthPNOor6EwKbn3R7iZVmY/edit"
    ]
)
ws3.column_dimensions["A"].width = 110

wb.save(out)
print("Saved", out)
print("Client:", len(client_posts), "Pro:", len(pro_posts))
